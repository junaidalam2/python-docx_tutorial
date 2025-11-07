import docx

doc = docx.Document('pyDoc.docx')

print(doc.paragraphs[0].text)  
print(doc.paragraphs[2].text)  
print(doc.paragraphs[0].runs[0].text)
print(doc.paragraphs[0].runs[1].text)
print(doc.paragraphs[0].runs[2].text)
print(doc.paragraphs[0].runs[3].text)

doc2 = docx.Document()
doc2.add_paragraph('Hello, World!', style='Heading 1')
paraObject = doc2.add_paragraph('This is a test paragraph.')

paraObject.add_run(' This is an added run.').bold = True

doc2.save('test.docx')